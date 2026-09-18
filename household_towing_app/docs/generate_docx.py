import os
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    import subprocess
    subprocess.check_call(["pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Add Title
title = doc.add_heading('DualServe System Data Dictionary', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

tables_data = [
    ("1. Users Data Table (users)", 
     "This table stores the core account credentials and basic profile information for every individual who uses the DualServe system. It acts as the central hub for authentication and role management, ensuring that customers, providers, and employees are properly identified before they can access their specific features.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Document ID, matches Firebase Auth UID."],
        ["name", "String", "Text", "100-255", "Full name of the user."],
        ["email", "String", "Email format", "255", "Contact email address."],
        ["phone", "String", "Numeric/String", "11-15", "Contact phone number."],
        ["role", "String", "Enum", "Varies", "User role: customer, provider, Employee."],
        ["createdAt", "Timestamp", "ISO 8601", "—", "Date and time the account was created."],
        ["profileImageUrl", "String", "URL", "255", "URL to the user's avatar."]
    ]),
    ("2. Providers Data Table (providers)", 
     "This table holds detailed business information for service companies (Towing agencies and Household Service providers). It extends the basic user profile by storing their company name, specialized skills, availability status, and public reputation metrics like ratings and reviews.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Matches the users ID."],
        ["specialty", "String", "Text", "50-100", "Primary skill or service type offered."],
        ["rating", "Double", "Decimal", "3", "Average star rating (1.0 to 5.0)."],
        ["totalReviews", "Integer", "Numeric", "Varies", "Total number of customer reviews received."],
        ["isAvailable", "Boolean", "True/False", "1", "Toggle for whether the provider is active."],
        ["companyName", "String", "Text", "100-255", "Official name of the towing company."],
        ["isApproved", "Boolean", "True/False", "1", "Administrator verification status."]
    ]),
    ("3. Employees Data Table (Employees)", 
     "This table manages the individual workers, such as tow truck drivers or household cleaners, who are employed by a Provider. It tracks their current duty availability, assigned vehicle or equipment, and professional license details.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Matches the users ID."],
        ["providerId", "String (Ref)", "Alphanumeric", "Varies", "The provider agency they work under."],
        ["vehicleType", "String", "Text", "50", "Type of vehicle assigned."],
        ["licenseNumber", "String", "Alphanumeric", "50", "Driver's license or certification number."],
        ["isAvailable", "Boolean", "True/False", "1", "Current duty status of the employee."]
    ]),
    ("4. Bookings Data Table (bookings)", 
     "This table records the formal service requests made by Customers. It captures the high-level details of the job, including the type of service requested, the estimated price, and the current lifecycle status (pending, accepted, completed) of the request before it becomes an active task.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the booking."],
        ["customerId", "String (Ref)", "Alphanumeric", "Varies", "UID of the customer."],
        ["assignedProviderId", "String (Ref)", "Alphanumeric", "Varies", "UID of the provider handling the booking."],
        ["serviceType", "String", "Enum", "50", "Type of service (e.g., Towing, Household)."],
        ["status", "String", "Enum", "20", "pending, accepted, rejected, completed, cancelled."],
        ["estimatedCost", "Double", "Decimal", "Varies", "The calculated price shown to the customer."],
        ["createdAt", "Timestamp", "ISO 8601", "—", "Date and time the booking was initiated."]
    ]),
    ("5. Tasks Data Table (tasks)", 
     "This table tracks the real-time execution of an accepted booking. It manages the dispatching process by linking a specific Employee to the job and providing the exact GPS coordinates and physical address needed for live map tracking and navigation.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the task."],
        ["bookingId", "String (Ref)", "Alphanumeric", "Varies", "Link to the original booking record."],
        ["assignedDriverId", "String (Ref)", "Alphanumeric", "Varies", "Employee/Driver dispatched to the task."],
        ["status", "String", "Enum", "20", "assigned, inProgress, completed, cancelled."],
        ["location", "String", "Text", "255", "Physical address of the service."],
        ["latitude", "Double", "Coordinate", "Varies", "GPS latitude for map routing."],
        ["longitude", "Double", "Coordinate", "Varies", "GPS longitude for map routing."]
    ]),
    ("6. Assets Data Table (assets)", 
     "This table inventories the physical equipment (like tow trucks, cleaning tools, or specialized gear) owned by a Provider. It monitors the current operational status of the equipment and tracks which Employee is actively using it in the field.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the asset."],
        ["ownerId", "String (Ref)", "Alphanumeric", "Varies", "Provider ID who owns the asset."],
        ["type", "String", "Enum", "20", "vehicle, tool, equipment."],
        ["category", "String", "Text", "50", "Specific classification (e.g., Tow Truck)."],
        ["status", "String", "Enum", "20", "active, maintenance, inactive, inUse."],
        ["assignedTo", "String (Ref)", "Alphanumeric", "Varies", "Employee ID currently using the asset."]
    ]),
    ("7. Transactions Data Table (transactions)", 
     "This table serves as the financial ledger for the system. It records the final, confirmed amount charged to a customer upon the completion of a booking, ensuring accurate financial tracking between the Customer and the Provider.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the transaction."],
        ["bookingId", "String (Ref)", "Alphanumeric", "Varies", "Associated booking reference."],
        ["customerId", "String (Ref)", "Alphanumeric", "Varies", "UID of the paying customer."],
        ["providerId", "String (Ref)", "Alphanumeric", "Varies", "UID of the receiving provider."],
        ["amount", "Double", "Decimal", "Varies", "Final price charged to the customer."],
        ["paymentStatus", "String", "Enum", "20", "pending, recorded, paid."]
    ]),
    ("8. Reviews Data Table (reviews)", 
     "This table stores the post-service feedback submitted by Customers. It maintains accountability and quality control by securely logging the star rating and written comments given to a specific Provider after a task is finished.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the review."],
        ["providerId", "String (Ref)", "Alphanumeric", "Varies", "UID of the reviewed provider."],
        ["customerId", "String (Ref)", "Alphanumeric", "Varies", "UID of the reviewer."],
        ["rating", "Double", "Decimal", "3", "Given star rating (1.0 to 5.0)."],
        ["comment", "String", "Text", "1000", "Feedback text from the customer."]
    ]),
    ("9. Messages Data Table (messages)", 
     "This table handles all in-app communication. It archives the real-time chat messages exchanged between Customers, Providers, and Employees during an active booking, ensuring that all conversations are timestamped and tied to a specific job.",
     [
        ["FIELD NAME", "DATA TYPE", "FORMAT", "LENGTH", "DESCRIPTION"],
        ["id", "String (ID)", "Alphanumeric", "Varies", "Unique identifier for the message."],
        ["bookingId", "String (Ref)", "Alphanumeric", "Varies", "Chat room/booking context ID."],
        ["senderId", "String (Ref)", "Alphanumeric", "Varies", "UID of the user who sent the message."],
        ["receiverId", "String (Ref)", "Alphanumeric", "Varies", "UID of the user receiving the message."],
        ["text", "String", "Text", "1000", "The actual chat content."],
        ["timestamp", "Timestamp", "ISO 8601", "—", "Date and time the message was sent."]
    ])
]

for title_text, description_text, rows in tables_data:
    # Add Heading
    heading = doc.add_heading(title_text, level=2)
    
    # Add Description
    desc = doc.add_paragraph(description_text)
    
    # Add Table
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    # Set header row
    hdr_cells = table.rows[0].cells
    for i, col_name in enumerate(rows[0]):
        hdr_cells[i].text = col_name
        # Make bold
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    
    # Add data rows
    for row_data in rows[1:]:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            
    doc.add_paragraph() # Add space between tables

out_path = r"c:\code\DUALSERVE\household_towing_app\docs\Data_Dictionary.docx"
doc.save(out_path)
print("SUCCESS")
